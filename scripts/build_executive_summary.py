"""Build the 1-2 slide executive-summary PPTX from tracked artifacts.

Outputs reports/output/executive_summary.pptx — the brief-mandated stakeholder
deck. Two slides: (1) the big picture with leaderboard chart, (2) operational
implications + future work.

Run:
    python scripts/build_executive_summary.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "reports" / "figures"
OUT = ROOT / "reports" / "output" / "executive_summary.pptx"

# Visual palette
NAVY = RGBColor(0x1A, 0x35, 0x52)
BLUE = RGBColor(0x1F, 0x77, 0xB4)
GREEN = RGBColor(0x2C, 0xA0, 0x2C)
RED = RGBColor(0xD6, 0x27, 0x28)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GREY = RGBColor(0xEE, 0xEE, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_text(frame, text: str, *, size: int, bold: bool = False,
              color: RGBColor = NAVY, align=PP_ALIGN.LEFT) -> None:
    frame.clear()
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_bullet(frame, text: str, *, size: int = 12, indent: int = 0,
                bold_lead: str | None = None) -> None:
    p = frame.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.level = indent
    if bold_lead:
        run = p.add_run()
        run.text = f"• {bold_lead} "
        run.font.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
    else:
        run = p.add_run()
        run.text = f"• {text}"
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"


def _add_bullets_to_frame(frame, items: list[tuple[str, str]], size: int = 12) -> None:
    """items is a list of (bold_lead, rest_of_text) tuples."""
    frame.clear()
    for i, (lead, rest) in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"• {lead} "
        run.font.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
        run2 = p.add_run()
        run2.text = rest
        run2.font.size = Pt(size)
        run2.font.color.rgb = NAVY
        run2.font.name = "Calibri"
        p.space_after = Pt(4)


def _add_colored_rectangle(slide, *, left, top, width, height, fill: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # blank

    # ── SLIDE 1: HEADLINE ─────────────────────────────────────────
    s = prs.slides.add_slide(blank_layout)

    # Top color band
    _add_colored_rectangle(s, left=Cm(0), top=Cm(0),
                           width=prs.slide_width, height=Cm(1.7), fill=NAVY)
    band_text = s.shapes.add_textbox(Cm(0.6), Cm(0.25), Cm(33), Cm(1.3)).text_frame
    band_text.margin_left = band_text.margin_right = Cm(0)
    band_text.margin_top = band_text.margin_bottom = Cm(0)
    p = band_text.paragraphs[0]
    r = p.add_run()
    r.text = "Forecasting Hourly Household Energy Consumption"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = "Calibri"

    p2 = band_text.add_paragraph()
    r = p2.add_run()
    r.text = ("LUISS Artificial Intelligence Techniques  •  Industry partner: Enel Global ICT  "
              "•  Thierry Ishimwe, Linda Carla Zorzoli, Mariavittoria Giurato, Cesar Dushimimana")
    r.font.size = Pt(10)
    r.font.color.rgb = WHITE
    r.font.name = "Calibri"

    # Hero metric on the left
    hero_box = s.shapes.add_textbox(Cm(0.6), Cm(2.2), Cm(7.5), Cm(4.5)).text_frame
    hero_box.margin_left = hero_box.margin_right = Cm(0.1)
    p = hero_box.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "MAE = 0.336 kW"
    r.font.size = Pt(36)
    r.font.bold = True
    r.font.color.rgb = BLUE
    r.font.name = "Calibri"
    p = hero_box.add_paragraph()
    r = p.add_run()
    r.text = "Cross-fold mean, XGBoost (leakage-safe pipeline)"
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = GREY
    r.font.name = "Calibri"
    p = hero_box.add_paragraph()
    p.space_before = Pt(10)
    r = p.add_run()
    r.text = ("Four model families (SARIMA, XGBoost, LSTM, GRU) compared on a uniform "
              "rolling-origin protocol with five metrics. Same-timestamp UCI columns "
              "(Voltage, Global_intensity, sub-metering) are mechanically excluded from "
              "the feature matrix; short-gap imputation is past-only. The pipeline ships "
              "with two regression tests guarding both guarantees.")
    r.font.size = Pt(10)
    r.font.color.rgb = NAVY
    r.font.name = "Calibri"

    # Leaderboard image on the right
    s.shapes.add_picture(
        str(FIGURES / "leaderboard_bar.png"),
        left=Cm(20.5), top=Cm(2.0),
        width=Cm(12.5), height=Cm(7.5),
    )

    # Headline findings underneath the hero block
    fbox = s.shapes.add_textbox(Cm(0.6), Cm(11.0), Cm(19.0), Cm(7.5)).text_frame
    fbox.word_wrap = True
    findings_1 = [
        ("Top three are a statistical tie.",
         "XGBoost 0.336, GRU 0.339, LSTM 0.341, SARIMA 0.345. Diebold–Mariano: "
         "XGBoost vs SARIMA (p=0.022) is the only significant pair at α=0.05."),
        ("GRU is competitive with LSTM at lower complexity.",
         "Same MAE within noise; the simpler GRU cell captures sequential structure "
         "as well or better at this resolution."),
        ("Weather augmentation validates the future-work priority.",
         "Adding 4 Paris-Montsouris weather columns to XGBoost: −0.77% MAE, "
         "concentrated on heating-season folds."),
        ("Leakage-safe by construction.",
         "Registry blocks same-timestamp components of the target; past-only "
         "imputation; two regression tests guard the policy."),
    ]
    _add_bullets_to_frame(fbox, findings_1, size=11)

    # ── SLIDE 2: OPERATIONAL IMPLICATIONS + FUTURE WORK ───────────
    s = prs.slides.add_slide(blank_layout)
    _add_colored_rectangle(s, left=Cm(0), top=Cm(0),
                           width=prs.slide_width, height=Cm(1.7), fill=NAVY)
    band_text = s.shapes.add_textbox(Cm(0.6), Cm(0.4), Cm(33), Cm(1.0)).text_frame
    p = band_text.paragraphs[0]
    r = p.add_run()
    r.text = "Where Models Win, Where They Lose, and What's Next"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = "Calibri"

    # Per-segment image on the left half
    s.shapes.add_picture(
        str(FIGURES / "segment_mae_by_dimension.png"),
        left=Cm(0.6), top=Cm(2.0),
        width=Cm(16.0), height=Cm(11.5),
    )

    # Caption under the figure
    cap = s.shapes.add_textbox(Cm(0.6), Cm(13.6), Cm(16.0), Cm(1.0)).text_frame
    p = cap.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = ("MAE by hour-of-day, day-of-week, and validation month. "
              "Peak hours and October heating-ramp dominate the error budget.")
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = GREY
    r.font.name = "Calibri"

    # Right column: operational implications + future work
    # Box 1 - Operational implications
    impl_title = s.shapes.add_textbox(Cm(17.5), Cm(2.0), Cm(15.5), Cm(1.0)).text_frame
    p = impl_title.paragraphs[0]
    r = p.add_run()
    r.text = "Operational implications"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = "Calibri"

    impl_box = s.shapes.add_textbox(Cm(17.5), Cm(3.0), Cm(15.5), Cm(6.5)).text_frame
    impl_box.word_wrap = True
    implications = [
        ("Peak-hour MAE is ~2.5× trough-hour MAE.",
         "XGBoost wins quiet hours; GRU wins peaks. For peak-cost-weighted operations, "
         "GRU may be the better deployment choice."),
        ("August easiest, October hardest.",
         "August = grandes vacances (empty household, flat load). October = regulated "
         "période de chauffe (heating ramp, MAE ~2× higher)."),
        ("Stacking does NOT break the 3% cluster.",
         "Base-model errors are too correlated. The bottleneck is input information, "
         "not model variety."),
    ]
    _add_bullets_to_frame(impl_box, implications, size=11)

    # Box 2 - Future work
    fw_title = s.shapes.add_textbox(Cm(17.5), Cm(10.0), Cm(15.5), Cm(1.0)).text_frame
    p = fw_title.paragraphs[0]
    r = p.add_run()
    r.text = "Future work"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = "Calibri"

    fw_box = s.shapes.add_textbox(Cm(17.5), Cm(11.0), Cm(15.5), Cm(7.0)).text_frame
    fw_box.word_wrap = True
    future = [
        ("Extend weather to SARIMA / LSTM / GRU.",
         "The −0.77% gain on XGBoost is a lower bound. ~70 min DL retrain on the "
         "already-built augmented matrix."),
        ("24-step-ahead day-ahead forecasting.",
         "Current results are 1-step-ahead at hourly resolution. Production at Enel "
         "uses 24-step horizons made once per day from a fixed cutoff."),
    ]
    _add_bullets_to_frame(fw_box, future, size=11)

    # Footer line on both slides
    for slide in prs.slides:
        footer = slide.shapes.add_textbox(Cm(0.6), Cm(18.2), Cm(32), Cm(0.6)).text_frame
        p = footer.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = ("github.com/ThierryIshimwe/Energy-Forecasting  •  "
                  "Brief: reports/Project_Work_Luiss_AI_Techniques_200326.pdf  •  "
                  "Full analysis: notebooks/00_main_deliverable.ipynb")
        r.font.size = Pt(8)
        r.font.italic = True
        r.font.color.rgb = GREY
        r.font.name = "Calibri"

    prs.save(str(OUT))
    print(f"Wrote {OUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
